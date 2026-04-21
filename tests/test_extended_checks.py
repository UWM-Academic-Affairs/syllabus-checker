"""
Tests targeting previously uncovered SyllabusChecker check methods:
  check_heading_hierarchy()
  check_layout_tables()
  check_empty_table_rows_columns()
  check_table_merged_cells()
  check_table_usage()
  check_manual_alignment()
  check_table_scope_declarations()
  analyze_sections_with_llm() — LLM code path
"""

import os
import json
import pytest
from docx import Document
from unittest.mock import MagicMock

from syllabus_checker import SyllabusChecker


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

def _make_checker(tmp_path_factory, name, builder_fn):
    d = tmp_path_factory.mktemp(name)
    path = str(d / f'{name}.docx')
    builder_fn(path)
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    template = os.path.join(root, 'assets', 'Uniform-Syllabus-Template-1.docx')
    return SyllabusChecker(template, path)


def _llm_response(content: str) -> MagicMock:
    resp = MagicMock()
    resp.choices = [MagicMock()]
    resp.choices[0].message.content = content
    return resp


# ---------------------------------------------------------------------------
# check_heading_hierarchy
# ---------------------------------------------------------------------------

class TestCheckHeadingHierarchy:

    @pytest.fixture(scope='class')
    def checker_all_caps_heading(self, tmp_path_factory):
        def build(path):
            doc = Document()
            doc.add_heading('GRADING SCALE', level=1)
            doc.add_paragraph('A 93-100, B 83-92.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'caps_heading', build)

    @pytest.fixture(scope='class')
    def checker_long_heading(self, tmp_path_factory):
        def build(path):
            doc = Document()
            long_text = 'This Is An Extremely Long Heading That Exceeds One Hundred Characters And Should Be Flagged By The Checker'
            doc.add_heading(long_text, level=1)
            doc.add_paragraph('Content.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'long_heading', build)

    @pytest.fixture(scope='class')
    def checker_skipped_level(self, tmp_path_factory):
        """H1 directly followed by H3 — skips H2."""
        def build(path):
            doc = Document()
            doc.add_heading('Major Section', level=1)
            doc.add_heading('Sub-subsection', level=3)
            doc.add_paragraph('Content.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'skipped_level', build)

    @pytest.fixture(scope='class')
    def checker_multiple_h1(self, tmp_path_factory):
        def build(path):
            doc = Document()
            doc.add_heading('Course Overview', level=1)
            doc.add_paragraph('Content.')
            doc.add_heading('Course Policies', level=1)
            doc.add_paragraph('Content.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'multiple_h1', build)

    def test_all_caps_heading_flagged(self, checker_all_caps_heading):
        issues = checker_all_caps_heading.check_heading_hierarchy()
        types = [i.issue_type for i in issues]
        assert 'ALL_CAPS_HEADING' in types

    def test_long_heading_flagged(self, checker_long_heading):
        issues = checker_long_heading.check_heading_hierarchy()
        types = [i.issue_type for i in issues]
        assert 'LONG_HEADING' in types

    def test_skipped_level_flagged(self, checker_skipped_level):
        issues = checker_skipped_level.check_heading_hierarchy()
        types = [i.issue_type for i in issues]
        # H1→H3 skip should suggest "SHOULD BE HEADING 2"
        assert any('HEADING' in t for t in types)

    def test_multiple_h1s_flagged(self, checker_multiple_h1):
        issues = checker_multiple_h1.check_heading_hierarchy()
        types = [i.issue_type for i in issues]
        assert 'MULTIPLE_H1' in types

    def test_returns_list(self, checker_all_caps_heading):
        assert isinstance(checker_all_caps_heading.check_heading_hierarchy(), list)


# ---------------------------------------------------------------------------
# check_layout_tables
# ---------------------------------------------------------------------------

class TestCheckLayoutTables:

    @pytest.fixture(scope='class')
    def checker_single_col(self, tmp_path_factory):
        """Single-column table — classic layout table pattern."""
        def build(path):
            doc = Document()
            table = doc.add_table(rows=3, cols=1)
            table.rows[0].cells[0].text = 'Row one content here'
            table.rows[1].cells[0].text = 'Row two content here'
            table.rows[2].cells[0].text = 'Row three content here'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'single_col', build)

    @pytest.fixture(scope='class')
    def checker_data_table(self, tmp_path_factory):
        """Well-formed data table with a bold header row — should not be a layout table."""
        def build(path):
            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            # Bold header
            for cell in table.rows[0].cells:
                run = cell.paragraphs[0].add_run('Header')
                run.bold = True
            table.rows[1].cells[0].text = 'Alpha'
            table.rows[1].cells[1].text = 'Beta'
            table.rows[2].cells[0].text = 'Gamma'
            table.rows[2].cells[1].text = 'Delta'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'data_table', build)

    def test_single_col_table_flagged_as_layout(self, checker_single_col):
        issues = checker_single_col.check_layout_tables()
        types = [i.issue_type for i in issues]
        assert 'LAYOUT_TABLE' in types

    def test_data_table_not_flagged(self, checker_data_table):
        issues = checker_data_table.check_layout_tables()
        types = [i.issue_type for i in issues]
        assert 'LAYOUT_TABLE' not in types


# ---------------------------------------------------------------------------
# check_empty_table_rows_columns
# ---------------------------------------------------------------------------

class TestCheckEmptyTableRowsColumns:

    @pytest.fixture(scope='class')
    def checker_empty_row(self, tmp_path_factory):
        """Table with one completely empty row used for spacing."""
        def build(path):
            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            table.rows[0].cells[0].text = 'Header A'
            table.rows[0].cells[1].text = 'Header B'
            # Row 1 left empty deliberately
            table.rows[2].cells[0].text = 'Data'
            table.rows[2].cells[1].text = 'Data'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'empty_row', build)

    @pytest.fixture(scope='class')
    def checker_no_empty_rows(self, tmp_path_factory):
        def build(path):
            doc = Document()
            table = doc.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = 'A'
            table.rows[0].cells[1].text = 'B'
            table.rows[1].cells[0].text = 'C'
            table.rows[1].cells[1].text = 'D'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'no_empty_rows', build)

    def test_empty_row_flagged(self, checker_empty_row):
        issues = checker_empty_row.check_empty_table_rows_columns()
        types = [i.issue_type for i in issues]
        assert 'EMPTY_TABLE_ROW' in types

    def test_no_empty_rows_not_flagged(self, checker_no_empty_rows):
        issues = checker_no_empty_rows.check_empty_table_rows_columns()
        types = [i.issue_type for i in issues]
        assert 'EMPTY_TABLE_ROW' not in types
        assert 'EMPTY_TABLE_COLUMN' not in types


# ---------------------------------------------------------------------------
# check_table_merged_cells
# ---------------------------------------------------------------------------

class TestCheckTableMergedCells:

    @pytest.fixture(scope='class')
    def checker_merged(self, tmp_path_factory):
        """Table with cells merged horizontally."""
        def build(path):
            doc = Document()
            table = doc.add_table(rows=2, cols=3)
            table.rows[0].cells[0].text = 'Alpha'
            table.rows[0].cells[1].text = 'Beta'
            table.rows[0].cells[2].text = 'Gamma'
            # Merge first two cells of row 1 horizontally
            table.cell(1, 0).merge(table.cell(1, 1))
            table.rows[1].cells[0].text = 'Merged cell'
            table.rows[1].cells[2].text = 'Separate'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'merged', build)

    @pytest.fixture(scope='class')
    def checker_no_merge(self, tmp_path_factory):
        def build(path):
            doc = Document()
            table = doc.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = 'A'
            table.rows[0].cells[1].text = 'B'
            table.rows[1].cells[0].text = 'C'
            table.rows[1].cells[1].text = 'D'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'no_merge', build)

    def test_merged_cells_flagged(self, checker_merged):
        issues = checker_merged.check_table_merged_cells()
        types = [i.issue_type for i in issues]
        assert 'TABLE_MERGED_CELLS' in types

    def test_no_merge_not_flagged(self, checker_no_merge):
        issues = checker_no_merge.check_table_merged_cells()
        types = [i.issue_type for i in issues]
        assert 'TABLE_MERGED_CELLS' not in types


# ---------------------------------------------------------------------------
# check_table_usage
# ---------------------------------------------------------------------------

class TestCheckTableUsage:

    @pytest.fixture(scope='class')
    def checker_many_tables(self, tmp_path_factory):
        """Document with 6 tables — exceeds the 5-table threshold."""
        def build(path):
            doc = Document()
            for _ in range(6):
                table = doc.add_table(rows=2, cols=2)
                table.rows[0].cells[0].text = 'A'
                table.rows[0].cells[1].text = 'B'
                table.rows[1].cells[0].text = 'C'
                table.rows[1].cells[1].text = 'D'
                doc.add_paragraph()
            doc.save(path)
        return _make_checker(tmp_path_factory, 'many_tables', build)

    @pytest.fixture(scope='class')
    def checker_large_cell(self, tmp_path_factory):
        """Table with a cell containing >500 characters."""
        def build(path):
            doc = Document()
            table = doc.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = 'Header'
            table.rows[0].cells[1].text = 'Value'
            table.rows[1].cells[0].text = 'A' * 501
            table.rows[1].cells[1].text = 'Short'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'large_cell', build)

    def test_many_tables_returns_warning(self, checker_many_tables):
        result = checker_many_tables.check_table_usage()
        assert isinstance(result['issues'], list)
        assert len(result['issues']) > 0
        assert any('table' in w.lower() for w in result['issues'])

    def test_large_cell_returns_warning(self, checker_large_cell):
        result = checker_large_cell.check_table_usage()
        assert len(result['issues']) > 0

    def test_returns_dict_with_count(self, checker_many_tables):
        result = checker_many_tables.check_table_usage()
        assert 'count' in result
        assert result['count'] == 6


# ---------------------------------------------------------------------------
# check_manual_alignment
# ---------------------------------------------------------------------------

class TestCheckManualAlignment:

    @pytest.fixture(scope='class')
    def checker_tab_aligned(self, tmp_path_factory):
        """Paragraph using a tab between two substantial pieces of text."""
        def build(path):
            doc = Document()
            # Need run_all_checks() first so the method can read self.issues
            doc.add_paragraph('Office Hours:\tMonday 2-4pm and Wednesday 1-3pm')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'tab_aligned', build)

    @pytest.fixture(scope='class')
    def checker_space_aligned(self, tmp_path_factory):
        """Paragraph using 4+ spaces for columnar alignment."""
        def build(path):
            doc = Document()
            doc.add_paragraph('Grade    Points    Percentage')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'space_aligned', build)

    @pytest.fixture(scope='class')
    def checker_normal_text(self, tmp_path_factory):
        def build(path):
            doc = Document()
            doc.add_paragraph('This is a completely normal paragraph with no special alignment.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'normal_text_ma', build)

    def test_tab_aligned_text_flagged(self, checker_tab_aligned):
        # check_manual_alignment reads self.issues; initialize to empty
        checker_tab_aligned.issues = []
        issues = checker_tab_aligned.check_manual_alignment()
        types = [i.issue_type for i in issues]
        assert 'PSEUDO_TABLE' in types

    def test_space_aligned_text_flagged(self, checker_space_aligned):
        checker_space_aligned.issues = []
        issues = checker_space_aligned.check_manual_alignment()
        types = [i.issue_type for i in issues]
        assert 'PSEUDO_TABLE' in types

    def test_normal_text_not_flagged(self, checker_normal_text):
        checker_normal_text.issues = []
        issues = checker_normal_text.check_manual_alignment()
        types = [i.issue_type for i in issues]
        assert 'PSEUDO_TABLE' not in types


# ---------------------------------------------------------------------------
# analyze_sections_with_llm — LLM code path (requires mocked OpenAI)
# ---------------------------------------------------------------------------

class TestAnalyzeSectionsWithLlmPath:
    """Cover the LLM branch of analyze_sections_with_llm() (lines 1302–1424)."""

    LLM_JSON = json.dumps({
        'found': [
            {'type': 'Welcome Statement', 'quote': 'Welcome to this course!'},
            {'type': 'Prerequisites', 'quote': 'COMPSCI 101 is required.'},
        ],
        'not_found': ['Resources', 'Department Policies'],
    })

    @pytest.fixture(scope='class')
    def checker(self, tmp_path_factory, fixture_dir):
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template = os.path.join(root, 'assets', 'Uniform-Syllabus-Template-1.docx')
        # Use missing_sections.docx so pattern detection leaves many items for the LLM
        return SyllabusChecker(template, str(fixture_dir / 'missing_sections.docx'))

    def test_llm_found_items_appear_in_present(self, checker, monkeypatch, mocker):
        monkeypatch.setenv('OPENAI_API_KEY', 'test-key')
        mock_cls = mocker.patch('syllabus_checker.OpenAI')
        mock_cls.return_value.chat.completions.create.return_value = (
            _llm_response(self.LLM_JSON)
        )

        result = checker.analyze_sections_with_llm()

        assert 'present' in result
        assert 'missing' in result
        present_types = [p['content_type'] for p in result['present']]
        assert 'Welcome Statement' in present_types or 'Prerequisites' in present_types

    def test_llm_not_found_items_appear_in_missing(self, checker, monkeypatch, mocker):
        monkeypatch.setenv('OPENAI_API_KEY', 'test-key')
        mock_cls = mocker.patch('syllabus_checker.OpenAI')
        mock_cls.return_value.chat.completions.create.return_value = (
            _llm_response(self.LLM_JSON)
        )

        result = checker.analyze_sections_with_llm()

        missing_types = [m['content_type'] for m in result['missing']]
        # Resources or Department Policies should be in missing (not found by LLM either)
        assert any(t in missing_types for t in ['Resources', 'Department Policies'])

    def test_llm_exception_falls_back_gracefully(self, checker, monkeypatch, mocker):
        """If the LLM call raises, remaining items are marked as needing review."""
        monkeypatch.setenv('OPENAI_API_KEY', 'test-key')
        mock_cls = mocker.patch('syllabus_checker.OpenAI')
        mock_cls.return_value.chat.completions.create.side_effect = (
            RuntimeError('connection refused')
        )

        result = checker.analyze_sections_with_llm()

        assert 'present' in result
        assert 'missing' in result

    def test_llm_malformed_json_falls_back_gracefully(self, checker, monkeypatch, mocker):
        """Malformed JSON from LLM falls into the exception handler."""
        monkeypatch.setenv('OPENAI_API_KEY', 'test-key')
        mock_cls = mocker.patch('syllabus_checker.OpenAI')
        mock_cls.return_value.chat.completions.create.return_value = (
            _llm_response('not valid JSON !!!')
        )

        result = checker.analyze_sections_with_llm()

        # Should not raise; returns dicts with present/missing keys
        assert 'present' in result
        assert 'missing' in result

    def test_returns_present_and_missing_keys(self, checker, monkeypatch, mocker):
        monkeypatch.setenv('OPENAI_API_KEY', 'test-key')
        mock_cls = mocker.patch('syllabus_checker.OpenAI')
        mock_cls.return_value.chat.completions.create.return_value = (
            _llm_response(self.LLM_JSON)
        )

        result = checker.analyze_sections_with_llm()

        assert isinstance(result['present'], list)
        assert isinstance(result['missing'], list)


# ---------------------------------------------------------------------------
# check_unstyled_headings
# ---------------------------------------------------------------------------

class TestCheckUnstyledHeadings:

    @pytest.fixture(scope='class')
    def checker_bold_heading(self, tmp_path_factory):
        """Short, fully-bold paragraph — should be flagged as an unstyled heading."""
        def build(path):
            doc = Document()
            para = doc.add_paragraph()
            run = para.add_run('Course Policies')
            run.bold = True
            doc.add_paragraph('Some descriptive content here that is not a heading.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'bold_unstyled_heading', build)

    @pytest.fixture(scope='class')
    def checker_allcaps_body(self, tmp_path_factory):
        """Short all-caps body paragraph — should be detected as unstyled heading."""
        def build(path):
            doc = Document()
            doc.add_paragraph('GRADING CRITERIA')
            doc.add_paragraph('Normal content follows here.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'allcaps_body', build)

    @pytest.fixture(scope='class')
    def checker_normal_paragraphs(self, tmp_path_factory):
        """Plain body text — no unstyled headings."""
        def build(path):
            doc = Document()
            doc.add_paragraph('This is a normal paragraph with regular text content.')
            doc.add_paragraph('Another sentence that ends with a period.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'normal_para_uh', build)

    def test_bold_paragraph_flagged(self, checker_bold_heading):
        issues = checker_bold_heading.check_unstyled_headings()
        types = [i.issue_type for i in issues]
        assert any('HEADING' in t for t in types)

    def test_allcaps_paragraph_flagged(self, checker_allcaps_body):
        issues = checker_allcaps_body.check_unstyled_headings()
        types = [i.issue_type for i in issues]
        assert any('HEADING' in t for t in types)

    def test_normal_paragraphs_not_flagged(self, checker_normal_paragraphs):
        issues = checker_normal_paragraphs.check_unstyled_headings()
        types = [i.issue_type for i in issues]
        assert not any('HEADING' in t for t in types)

    def test_returns_list(self, checker_bold_heading):
        assert isinstance(checker_bold_heading.check_unstyled_headings(), list)


# ---------------------------------------------------------------------------
# check_inconsistent_fonts
# ---------------------------------------------------------------------------

class TestCheckInconsistentFonts:

    @pytest.fixture(scope='class')
    def checker_many_fonts(self, tmp_path_factory):
        """Document using 3 fonts with significant secondary usage."""
        def build(path):
            doc = Document()
            # Primary font: Calibri (majority)
            for _ in range(8):
                para = doc.add_paragraph()
                run = para.add_run('This text uses Calibri for the main body content.')
                run.font.name = 'Calibri'
            # Secondary font: Times New Roman (>10% usage)
            for _ in range(2):
                para = doc.add_paragraph()
                run = para.add_run('This text uses Times New Roman as a secondary font.')
                run.font.name = 'Times New Roman'
            # Tertiary font: Arial (>10% usage)
            for _ in range(2):
                para = doc.add_paragraph()
                run = para.add_run('This text uses Arial as a third distinct font.')
                run.font.name = 'Arial'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'many_fonts', build)

    @pytest.fixture(scope='class')
    def checker_one_font(self, tmp_path_factory):
        """Document with a single consistent font throughout."""
        def build(path):
            doc = Document()
            for _ in range(5):
                para = doc.add_paragraph()
                run = para.add_run('Consistent Calibri text throughout the document.')
                run.font.name = 'Calibri'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'one_font', build)

    def test_many_fonts_flagged(self, checker_many_fonts):
        issues = checker_many_fonts.check_inconsistent_fonts()
        types = [i.issue_type for i in issues]
        assert 'INCONSISTENT_FONTS' in types

    def test_one_font_not_flagged(self, checker_one_font):
        issues = checker_one_font.check_inconsistent_fonts()
        types = [i.issue_type for i in issues]
        assert 'INCONSISTENT_FONTS' not in types

    def test_returns_list(self, checker_one_font):
        assert isinstance(checker_one_font.check_inconsistent_fonts(), list)


# ---------------------------------------------------------------------------
# check_unstyled_links
# ---------------------------------------------------------------------------

class TestCheckUnstyledLinks:

    @pytest.fixture(scope='class')
    def checker_no_links(self, tmp_path_factory):
        """Document with no hyperlinks — should return empty list."""
        def build(path):
            doc = Document()
            doc.add_paragraph('No links in this document, just plain text.')
            doc.save(path)
        return _make_checker(tmp_path_factory, 'no_links', build)

    def test_no_links_returns_empty_list(self, checker_no_links):
        issues = checker_no_links.check_unstyled_links()
        assert issues == []

    def test_returns_list(self, checker_no_links):
        assert isinstance(checker_no_links.check_unstyled_links(), list)


# ---------------------------------------------------------------------------
# check_table_scope_declarations
# ---------------------------------------------------------------------------

class TestCheckTableScopeDeclarations:

    @pytest.fixture(scope='class')
    def checker_bold_header_no_xml(self, tmp_path_factory):
        """
        Table where the first row is bold (heuristic header) but does NOT have
        the w:tblHeader XML element — should be flagged TABLE_MISSING_SCOPE.
        """
        def build(path):
            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            for cell in table.rows[0].cells:
                run = cell.paragraphs[0].add_run('Header')
                run.bold = True
            table.rows[1].cells[0].text = 'Data A'
            table.rows[1].cells[1].text = 'Data B'
            table.rows[2].cells[0].text = 'Data C'
            table.rows[2].cells[1].text = 'Data D'
            doc.save(path)
        return _make_checker(tmp_path_factory, 'bold_no_xml', build)

    def test_bold_header_without_tblHeader_flagged(self, checker_bold_header_no_xml):
        issues = checker_bold_header_no_xml.check_table_scope_declarations()
        types = [i.issue_type for i in issues]
        assert 'TABLE_MISSING_SCOPE' in types

    def test_returns_list(self, checker_bold_header_no_xml):
        assert isinstance(checker_bold_header_no_xml.check_table_scope_declarations(), list)
