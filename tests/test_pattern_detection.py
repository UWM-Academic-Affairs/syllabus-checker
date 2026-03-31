"""
Unit tests for SyllabusChecker pattern-detection methods:
  _detect_content_patterns()
  _find_heading_for_content()
  _extract_headings_with_positions()
  check_missing_sections()
  _get_section_description()

Most methods accept plain strings, so the tests pass crafted text directly
rather than building complex DOCX fixtures for every case.
"""

import pytest
from docx import Document
from syllabus_checker import SyllabusChecker


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(scope='session')
def checker(template_path, fixture_dir):
    """SyllabusChecker backed by minimal_valid.docx (all sections present)."""
    return SyllabusChecker(template_path, str(fixture_dir / 'minimal_valid.docx'))


@pytest.fixture(scope='session')
def checker_missing(template_path, fixture_dir):
    """SyllabusChecker backed by missing_sections.docx (most sections absent)."""
    return SyllabusChecker(template_path, str(fixture_dir / 'missing_sections.docx'))


@pytest.fixture(scope='session')
def checker_table_headings(template_path, tmp_path_factory):
    """SyllabusChecker whose target DOCX has both a top-level AND a table-cell heading."""
    d = tmp_path_factory.mktemp('table_hdg')
    path = str(d / 'table_headings.docx')
    doc = Document()
    # Top-level heading (in doc.paragraphs → location: 'top-level')
    doc.add_heading('Instructor Name & Contact', level=1)
    doc.add_paragraph('Dr. Smith  jsmith@uwm.edu')
    # Table-cell heading (in a table cell → location: 'Table 1, Row 1, Cell 1')
    table = doc.add_table(rows=1, cols=2)
    cell = table.rows[0].cells[0]
    para = cell.paragraphs[0]
    para.text = 'Note'
    para.style = doc.styles['Heading 2']
    doc.save(path)
    return SyllabusChecker(template_path, path)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _h(text, position, *, is_word_heading=False, is_top_level=False):
    """Minimal heading dict for _find_heading_for_content tests."""
    return {
        'text': text,
        'position': position,
        'standard_name': None,
        'is_word_heading': is_word_heading,
        'is_top_level': is_top_level,
    }


# ---------------------------------------------------------------------------
# _detect_content_patterns
# ---------------------------------------------------------------------------

class TestDetectContentPatterns:

    def test_email_detects_instructor_section(self, checker):
        text = "Instructor Name & Contact\nDr. Smith\njsmith@uwm.edu\nOffice: EMS 123\n"
        result = checker._detect_content_patterns(text)
        assert 'Instructor' in result
        assert 'jsmith@uwm.edu' in result['Instructor']['found']

    def test_phone_number_also_detects_instructor(self, checker):
        text = "Instructor Name & Contact\nDr. Smith\n(414) 229-4567\n"
        result = checker._detect_content_patterns(text)
        assert 'Instructor' in result

    def test_inline_letter_grade_detects_grading_scale(self, checker):
        # Standard format: A: 93 or A=93
        text = "Grading Scale\nA: 93-100\nB: 83-92\nC: 73-82\n"
        result = checker._detect_content_patterns(text)
        assert 'Grading Scale' in result

    def test_table_cell_format_detects_grading_scale(self, checker):
        # Table-cell format: letter and number land in separate cells,
        # joined by whitespace in the extracted text (e.g. "A 93")
        text = "Grading Scale\nA 93\nB 83\nC 73\n"
        result = checker._detect_content_patterns(text)
        assert 'Grading Scale' in result

    def test_week_number_pattern_detects_calendar(self, checker):
        text = "Course Schedule\nWeek 1: Introduction\nWeek 2: Data Types\n"
        result = checker._detect_content_patterns(text)
        assert 'Calendar' in result

    def test_month_date_pattern_detects_calendar(self, checker):
        text = "Important Dates\nJanuary 15: First class\nFebruary 10: Midterm\n"
        result = checker._detect_content_patterns(text)
        assert 'Calendar' in result

    def test_placeholder_markers_suppress_heading_rename_suggestion(self, checker):
        """
        When the found snippet contains a template placeholder marker (< - ... - >),
        is_placeholder() flips heading_ok to True so no rename suggestion fires —
        even though there is no proper heading context.
        """
        # The grade pattern fires on "A: 93", but the snippet includes a placeholder
        # marker so heading_ok must be suppressed to True.
        text = "A: 93 < - sample grade placeholder - >\n"
        result = checker._detect_content_patterns(text)
        assert 'Grading Scale' in result
        assert result['Grading Scale']['heading_ok'] is True


# ---------------------------------------------------------------------------
# _find_heading_for_content
# ---------------------------------------------------------------------------

class TestFindHeadingForContent:

    def test_pseudo_heading_closer_than_top_level_wins(self, checker):
        """A pseudo-heading near the content beats a distant document title."""
        headings = [
            _h('Course Syllabus', 0,   is_word_heading=True, is_top_level=True),
            _h('Grading:',        100),   # pseudo-heading, closer to content
        ]
        result = checker._find_heading_for_content(200, headings)
        assert result['text'] == 'Grading:'

    def test_top_level_returned_when_no_closer_pseudo(self, checker):
        headings = [
            _h('Grading Scale', 50, is_word_heading=True, is_top_level=True),
        ]
        result = checker._find_heading_for_content(200, headings)
        assert result['text'] == 'Grading Scale'

    def test_table_cell_word_heading_does_not_override_top_level(self, checker):
        """
        A table-cell Word heading (is_word_heading=True, is_top_level=False) that
        is closer to the content must NOT beat the enclosing top-level heading.
        """
        headings = [
            _h('Grading Scale', 0, is_word_heading=True, is_top_level=True),
            _h('Note',         50, is_word_heading=True, is_top_level=False),
        ]
        result = checker._find_heading_for_content(100, headings)
        assert result['text'] == 'Grading Scale'

    def test_no_headings_returns_none(self, checker):
        assert checker._find_heading_for_content(100, []) is None

    def test_heading_after_content_is_ignored(self, checker):
        """Headings that come AFTER the content position must be ignored."""
        headings = [
            _h('Future Section', 200, is_word_heading=True, is_top_level=True),
        ]
        assert checker._find_heading_for_content(100, headings) is None


# ---------------------------------------------------------------------------
# _extract_headings_with_positions
# ---------------------------------------------------------------------------

class TestExtractHeadingsWithPositions:

    # Text used by both tests — contains both heading strings so the
    # method can find them by scanning full_text.
    _TEXT = (
        "Instructor Name & Contact\n"
        "Dr. Smith  jsmith@uwm.edu\n"
        "Note\n"
        "Some table cell content here\n"
    )

    def test_top_level_heading_has_is_top_level_true(self, checker_table_headings):
        headings = checker_table_headings._extract_headings_with_positions(self._TEXT)
        word_headings = [h for h in headings if h.get('is_word_heading')]
        top_level_texts = [h['text'] for h in word_headings if h.get('is_top_level')]
        assert 'Instructor Name & Contact' in top_level_texts

    def test_table_cell_heading_has_is_top_level_false(self, checker_table_headings):
        headings = checker_table_headings._extract_headings_with_positions(self._TEXT)
        word_headings = [h for h in headings if h.get('is_word_heading')]
        non_top_texts = [h['text'] for h in word_headings if not h.get('is_top_level')]
        assert 'Note' in non_top_texts


# ---------------------------------------------------------------------------
# check_missing_sections
# ---------------------------------------------------------------------------

class TestCheckMissingSections:

    def test_minimal_syllabus_has_expected_missing_sections(self, checker_missing):
        missing = checker_missing.check_missing_sections()
        assert 'Instructor Name & Contact' in missing
        assert 'Welcome Statement' in missing
        assert 'Grading Scale' in missing

    def test_complete_syllabus_has_no_false_positives(self, checker):
        missing = checker.check_missing_sections()
        assert 'Instructor Name & Contact' not in missing
        assert 'Grading Scale' not in missing
        assert 'Calendar' not in missing
        assert 'Welcome Statement' not in missing

    def test_returns_list(self, checker):
        assert isinstance(checker.check_missing_sections(), list)


# ---------------------------------------------------------------------------
# _get_section_description
# ---------------------------------------------------------------------------

class TestGetSectionDescription:

    def test_grading_scale_includes_sample_grade_table_rows(self, checker):
        desc = checker._get_section_description('Grading Scale')
        # The template embeds a sample grade table; it should be rendered as
        # pipe-delimited rows (table_to_html format) containing letter grades.
        assert '|' in desc or any(g in desc for g in ['93', '90', 'A-', 'B+', 'A+'])

    def test_grading_scale_is_not_generic_fallback(self, checker):
        desc = checker._get_section_description('Grading Scale')
        generic = "This section helps students understand grading scale."
        # The real template description must be richer than the fallback
        assert desc != generic

    def test_returns_string(self, checker):
        assert isinstance(checker._get_section_description('Grading Scale'), str)

    def test_unknown_section_returns_non_empty_fallback(self, checker):
        desc = checker._get_section_description('Nonexistent Section XYZ')
        assert isinstance(desc, str)
        assert len(desc) > 0
