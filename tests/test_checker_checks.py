"""
Tests for SyllabusChecker check methods not covered elsewhere:
  run_all_checks()
  check_heading_structure()
  check_decorative_fonts()
  check_line_spacing()
  check_document_metadata()
"""

import os
import pytest
from docx import Document
from docx.shared import Pt

from syllabus_checker import SyllabusChecker


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _checker(tmp_path_factory, name, builder_fn):
    d = tmp_path_factory.mktemp(name)
    path = str(d / f'{name}.docx')
    builder_fn(path)
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    template = os.path.join(root, 'assets', 'Uniform-Syllabus-Template-1.docx')
    return SyllabusChecker(template, path)


# ---------------------------------------------------------------------------
# run_all_checks
# ---------------------------------------------------------------------------

class TestRunAllChecks:

    @pytest.fixture(scope='class')
    def checker(self, tmp_path_factory, fixture_dir):
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template = os.path.join(root, 'assets', 'Uniform-Syllabus-Template-1.docx')
        return SyllabusChecker(template, str(fixture_dir / 'minimal_valid.docx'))

    def test_populates_issues_list(self, checker):
        checker.run_all_checks()
        assert isinstance(checker.issues, list)

    def test_does_not_raise(self, checker):
        """run_all_checks() must complete without raising on a well-formed DOCX."""
        checker.run_all_checks()  # should not raise

    def test_issues_are_accessibility_issue_objects(self, checker):
        from syllabus_checker import AccessibilityIssue
        checker.run_all_checks()
        for issue in checker.issues:
            assert isinstance(issue, AccessibilityIssue)

    def test_each_issue_has_type_and_description(self, checker):
        checker.run_all_checks()
        for issue in checker.issues:
            assert issue.issue_type
            assert issue.description


# ---------------------------------------------------------------------------
# check_heading_structure
# ---------------------------------------------------------------------------

class TestCheckHeadingStructure:

    @pytest.fixture(scope='class')
    def checker_no_headings(self, tmp_path_factory):
        """DOCX with only plain paragraphs — no Heading styles at all."""
        def build(path):
            doc = Document()
            doc.add_paragraph('Just some plain text with no headings at all.')
            doc.add_paragraph('Another plain paragraph.')
            doc.save(path)
        return _checker(tmp_path_factory, 'no_headings', build)

    @pytest.fixture(scope='class')
    def checker_only_h1(self, tmp_path_factory):
        """DOCX with only Heading 1 — no sub-levels."""
        def build(path):
            doc = Document()
            doc.add_heading('Section One', level=1)
            doc.add_paragraph('Content under section one.')
            doc.add_heading('Section Two', level=1)
            doc.add_paragraph('Content under section two.')
            doc.save(path)
        return _checker(tmp_path_factory, 'only_h1', build)

    @pytest.fixture(scope='class')
    def checker_h2_first(self, tmp_path_factory):
        """DOCX that starts with Heading 2, not Heading 1."""
        def build(path):
            doc = Document()
            doc.add_heading('A Sub-Section', level=2)
            doc.add_paragraph('Content.')
            doc.save(path)
        return _checker(tmp_path_factory, 'h2_first', build)

    @pytest.fixture(scope='class')
    def checker_many_h1s(self, tmp_path_factory):
        """DOCX with >10 Heading 1s and no Heading 2s — hierarchy warning."""
        def build(path):
            doc = Document()
            for i in range(12):
                doc.add_heading(f'Major Section {i+1}', level=1)
                doc.add_paragraph(f'Content {i+1}.')
            doc.save(path)
        return _checker(tmp_path_factory, 'many_h1s', build)

    @pytest.fixture(scope='class')
    def checker_good_structure(self, tmp_path_factory, fixture_dir):
        """minimal_valid.docx has proper H1 structure."""
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template = os.path.join(root, 'assets', 'Uniform-Syllabus-Template-1.docx')
        return SyllabusChecker(template, str(fixture_dir / 'minimal_valid.docx'))

    def test_no_headings_returns_critical_warning(self, checker_no_headings):
        result = checker_no_headings.check_heading_structure()
        assert any('CRITICAL' in issue for issue in result['issues'])

    def test_no_headings_counts_dict_is_empty(self, checker_no_headings):
        result = checker_no_headings.check_heading_structure()
        assert result['counts'] == {}

    def test_only_h1_warns_limited_hierarchy(self, checker_only_h1):
        result = checker_only_h1.check_heading_structure()
        warnings = ' '.join(result['issues'])
        assert 'WARNING' in warnings or 'hierarchy' in warnings.lower()

    def test_h2_first_warns_no_h1(self, checker_h2_first):
        result = checker_h2_first.check_heading_structure()
        warnings = ' '.join(result['issues'])
        assert 'Heading 1' in warnings or 'heading 1' in warnings.lower()

    def test_many_h1s_warns_hierarchy(self, checker_many_h1s):
        result = checker_many_h1s.check_heading_structure()
        assert len(result['issues']) > 0

    def test_returns_dict_with_counts_and_issues_keys(self, checker_good_structure):
        result = checker_good_structure.check_heading_structure()
        assert 'counts' in result
        assert 'issues' in result

    def test_good_structure_has_h1_in_counts(self, checker_good_structure):
        result = checker_good_structure.check_heading_structure()
        assert 'Heading 1' in result['counts']


# ---------------------------------------------------------------------------
# check_decorative_fonts
# ---------------------------------------------------------------------------

class TestCheckDecorativeFonts:

    @pytest.fixture(scope='class')
    def checker_comic_sans(self, tmp_path_factory):
        """DOCX with a run using Comic Sans MS — a decorative font."""
        def build(path):
            doc = Document()
            para = doc.add_paragraph()
            run = para.add_run('This text uses a decorative font.')
            run.font.name = 'Comic Sans MS'
            doc.save(path)
        return _checker(tmp_path_factory, 'comic_sans', build)

    @pytest.fixture(scope='class')
    def checker_standard_font(self, tmp_path_factory):
        """DOCX with Calibri — a standard readable font."""
        def build(path):
            doc = Document()
            para = doc.add_paragraph()
            run = para.add_run('This text uses a standard font.')
            run.font.name = 'Calibri'
            doc.save(path)
        return _checker(tmp_path_factory, 'standard_font', build)

    def test_comic_sans_flagged(self, checker_comic_sans):
        issues = checker_comic_sans.check_decorative_fonts()
        types = [i.issue_type for i in issues]
        assert 'DECORATIVE_FONT' in types

    def test_standard_font_not_flagged(self, checker_standard_font):
        issues = checker_standard_font.check_decorative_fonts()
        types = [i.issue_type for i in issues]
        assert 'DECORATIVE_FONT' not in types


# ---------------------------------------------------------------------------
# check_line_spacing
# ---------------------------------------------------------------------------

class TestCheckLineSpacing:

    @pytest.fixture(scope='class')
    def checker_tight_spacing(self, tmp_path_factory):
        """DOCX paragraph with line spacing set to 0.8 (below 1.15 threshold)."""
        def build(path):
            from docx.shared import Pt as _Pt
            doc = Document()
            para = doc.add_paragraph('This paragraph has very tight line spacing that is hard to read.')
            para.paragraph_format.line_spacing = 0.8
            doc.save(path)
        return _checker(tmp_path_factory, 'tight_spacing', build)

    @pytest.fixture(scope='class')
    def checker_normal_spacing(self, tmp_path_factory):
        """DOCX paragraph with line spacing of 1.5 (well above threshold)."""
        def build(path):
            doc = Document()
            para = doc.add_paragraph('This paragraph has comfortable line spacing.')
            para.paragraph_format.line_spacing = 1.5
            doc.save(path)
        return _checker(tmp_path_factory, 'normal_spacing', build)

    def test_tight_spacing_flagged(self, checker_tight_spacing):
        issues = checker_tight_spacing.check_line_spacing()
        types = [i.issue_type for i in issues]
        assert 'LOW_LINE_SPACING' in types

    def test_normal_spacing_not_flagged(self, checker_normal_spacing):
        issues = checker_normal_spacing.check_line_spacing()
        types = [i.issue_type for i in issues]
        assert 'LOW_LINE_SPACING' not in types


# ---------------------------------------------------------------------------
# check_document_metadata
# ---------------------------------------------------------------------------

class TestCheckDocumentMetadata:

    @pytest.fixture(scope='class')
    def checker_no_title(self, tmp_path_factory):
        """DOCX whose core_properties.title is empty (the default)."""
        def build(path):
            doc = Document()
            doc.add_paragraph('A document without a title set in metadata.')
            # Leave doc.core_properties.title as the default empty string
            doc.save(path)
        return _checker(tmp_path_factory, 'no_title', build)

    @pytest.fixture(scope='class')
    def checker_with_title(self, tmp_path_factory):
        """DOCX with a non-empty title in core properties."""
        def build(path):
            doc = Document()
            doc.add_paragraph('A document with a proper title.')
            doc.core_properties.title = 'Introduction to Testing — Syllabus'
            doc.save(path)
        return _checker(tmp_path_factory, 'with_title', build)

    def test_missing_title_flagged(self, checker_no_title):
        issues = checker_no_title.check_document_metadata()
        types = [i.issue_type for i in issues]
        assert 'MISSING_TITLE' in types

    def test_present_title_not_flagged(self, checker_with_title):
        issues = checker_with_title.check_document_metadata()
        types = [i.issue_type for i in issues]
        assert 'MISSING_TITLE' not in types
