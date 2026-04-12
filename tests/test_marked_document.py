"""
Tests for SyllabusChecker.create_marked_document():
  - Produces a valid .docx file
  - Missing sections appear in the output document's appended section
  - No crash when missing_sections is empty
  - No crash when growth_mindset_analysis is absent or errored
"""

import os
import pytest
from docx import Document

from syllabus_checker import SyllabusChecker


# ---------------------------------------------------------------------------
# Shared fixture
# ---------------------------------------------------------------------------

@pytest.fixture(scope='module')
def checker(tmp_path_factory, fixture_dir):
    """SyllabusChecker backed by minimal_valid.docx."""
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    template = os.path.join(root, 'assets', 'Uniform-Syllabus-Template-1.docx')
    return SyllabusChecker(template, str(fixture_dir / 'minimal_valid.docx'))


@pytest.fixture
def out(tmp_path):
    """Return a path for the marked output file."""
    return str(tmp_path / 'marked.docx')


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestCreateMarkedDocument:

    def test_produces_valid_docx(self, checker, out):
        """Output file must be openable as a DOCX without error."""
        checker.create_marked_document(out)
        assert os.path.exists(out)
        doc = Document(out)           # raises if file is not a valid .docx
        assert len(doc.paragraphs) > 0

    def test_missing_sections_appear_in_output(self, checker, out):
        """Named missing sections should appear as text in the output document."""
        missing = ['Welcome Statement', 'Grading Scale']
        checker.create_marked_document(out, missing_sections=missing)
        doc = Document(out)
        all_text = '\n'.join(p.text for p in doc.paragraphs)
        assert 'Welcome Statement' in all_text
        assert 'Grading Scale' in all_text

    def test_missing_sections_banner_present(self, checker, out):
        """The 'MISSING SECTIONS' banner paragraph must appear when sections are missing."""
        checker.create_marked_document(out, missing_sections=['Prerequisites'])
        doc = Document(out)
        all_text = '\n'.join(p.text for p in doc.paragraphs)
        assert 'MISSING SECTIONS' in all_text

    def test_no_crash_when_missing_sections_empty(self, checker, out):
        """Passing an empty list must not raise and must produce a valid file."""
        checker.create_marked_document(out, missing_sections=[])
        assert os.path.exists(out)
        Document(out)  # must open without error

    def test_no_crash_when_missing_sections_omitted(self, checker, out):
        """Omitting missing_sections entirely (None default) must not raise."""
        checker.create_marked_document(out)
        assert os.path.exists(out)
        Document(out)

    def test_no_crash_when_growth_mindset_absent(self, checker, out):
        """Passing None for growth_mindset_analysis must not raise."""
        checker.create_marked_document(out, growth_mindset_analysis=None)
        assert os.path.exists(out)
        Document(out)

    def test_no_crash_when_growth_mindset_errored(self, checker, out):
        """A dict with status='error' must not raise (banner is skipped)."""
        checker.create_marked_document(
            out,
            growth_mindset_analysis={'status': 'error', 'message': 'API timeout'},
        )
        assert os.path.exists(out)
        doc = Document(out)
        all_text = '\n'.join(p.text for p in doc.paragraphs)
        # The growth-mindset banner must NOT appear for an errored analysis
        assert 'BELONGING & GROWTH MINDSET' not in all_text

    def test_growth_mindset_banner_appears_on_success(self, checker, out):
        """A dict with status='success' and 'analysis' key renders the banner."""
        checker.create_marked_document(
            out,
            growth_mindset_analysis={
                'status': 'success',
                'analysis': 'Consider adding a welcome message that emphasizes growth.',
            },
        )
        doc = Document(out)
        all_text = '\n'.join(p.text for p in doc.paragraphs)
        assert 'BELONGING & GROWTH MINDSET' in all_text
