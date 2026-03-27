"""
Pytest configuration and shared fixtures for the Syllabus Checker test suite.

DOCX fixtures are generated programmatically at test-session start (rather than
committed to the repo) because *.docx is listed in .gitignore.
"""

import os
import pytest
from unittest.mock import MagicMock

# Keep Flask in development mode so SESSION_COOKIE_SECURE stays off and the
# startup SECRET_KEY warning is suppressed during tests.
os.environ.setdefault('FLASK_ENV', 'development')
# Ensure no real OpenAI calls escape during tests.
os.environ.pop('OPENAI_API_KEY', None)

import analytics_db  # noqa: E402 — must come after env setup
from app import app as flask_app, limiter as flask_limiter  # noqa: E402


# ---------------------------------------------------------------------------
# DOCX fixture factories
# ---------------------------------------------------------------------------

def _build_minimal_valid_docx(path: str) -> None:
    """DOCX that satisfies every required section."""
    from docx import Document

    doc = Document()
    sections = [
        ('Instructor Name & Contact',
         'Dr. Jane Smith  •  jsmith@uwm.edu  •  Office: EMS 123'),
        ('Office Hours',
         'Monday 2–4 PM, Wednesday 1–3 PM, or by appointment.'),
        ('Course Title & Credits',
         'Introduction to Testing — 3 credits'),
        ('Meeting Times & Location',
         'MWF 10:00–10:50 AM, EMS 180'),
        ('Welcome Statement',
         'Welcome to this course! I am glad you are here.'),
        ('Course Objectives',
         'Students will learn to design and execute software tests.'),
        ('Student Learning Outcomes',
         'Upon completion students will be able to: '
         '(1) write unit tests, (2) debug failing tests.'),
        ('Assessment',
         'Students are assessed via quizzes, projects, and a final exam.'),
        ('Prerequisites',
         'COMPSCI 101 or equivalent programming experience required.'),
        ('Course Modality and Format',
         'In-person course meeting three times per week.'),
        ('Course Materials',
         'No required textbook. All materials posted on Canvas.'),
        ('Time Investment',
         'Expect approximately 9 hours per week outside of class.'),
        ('Assignment and Grading Policies',
         'Assignments are due at 11:59 PM. Late work loses 10% per day.'),
        ('Grading Scheme',
         'Quizzes 20 %  •  Projects 50 %  •  Final Exam 30 %'),
        ('Grading Scale',
         'A: 93–100  A-: 90–92  B+: 87–89  B: 83–86  B-: 80–82'),
        ('Course Policies',
         'Attendance is expected. Please arrive on time.'),
        ('Department Policies',
         'See the department website for all department-level policies.'),
        ('University Policies',
         'UWM Academic Integrity, Accessibility, and Title IX policies apply.'),
        ('Resources',
         'Writing Center, Tutoring Center, and Mental Health Services available.'),
        ('Important Dates',
         'Midterm: Week 8.  Final Exam: December 15.'),
        ('Instructor Statement of Support',
         'I want you to succeed. Please reach out early if you need help.'),
        ('Calendar',
         'Week 1: Introduction\nWeek 2: Unit Testing\nWeek 3: Integration Testing'),
    ]
    for heading_text, content in sections:
        doc.add_heading(heading_text, level=1)
        doc.add_paragraph(content)
    doc.save(path)


def _build_missing_sections_docx(path: str) -> None:
    """DOCX that is clearly missing most required sections."""
    from docx import Document

    doc = Document()
    doc.add_heading('Course Information', level=1)
    doc.add_paragraph('This is a very minimal syllabus with incomplete information.')
    doc.add_heading('Grading', level=1)
    doc.add_paragraph('You will be graded on participation.')
    doc.save(path)


def _build_malformed_docx(path: str) -> None:
    """Write random bytes — not a valid ZIP/DOCX file at all."""
    with open(path, 'wb') as fh:
        fh.write(b'This is NOT a valid DOCX file.\x00\x01\x02\x03')


# ---------------------------------------------------------------------------
# Session-scoped fixture directory (created once per test run)
# ---------------------------------------------------------------------------

@pytest.fixture(scope='session')
def fixture_dir(tmp_path_factory):
    """Return a Path containing all generated DOCX test fixtures."""
    d = tmp_path_factory.mktemp('fixtures')
    _build_minimal_valid_docx(str(d / 'minimal_valid.docx'))
    _build_missing_sections_docx(str(d / 'missing_sections.docx'))
    _build_malformed_docx(str(d / 'malformed.docx'))
    return d


@pytest.fixture(scope='session')
def template_path():
    """Absolute path to the bundled syllabus template (never swept)."""
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    p = os.path.join(root, 'assets', 'Uniform-Syllabus-Template-1.docx')
    assert os.path.exists(p), f'Template not found at {p}'
    return p


# ---------------------------------------------------------------------------
# Flask app + test client
# ---------------------------------------------------------------------------

@pytest.fixture
def app(tmp_path, monkeypatch):
    """Isolated Flask app instance with a temp DB, upload dir, and repo dir."""
    upload_dir = str(tmp_path / 'uploads')
    repo_dir = str(tmp_path / 'repo')
    db_path = str(tmp_path / 'test.db')
    os.makedirs(upload_dir, exist_ok=True)
    os.makedirs(repo_dir, exist_ok=True)

    # Redirect analytics DB to a fresh temp file for this test.
    # get_db_connection() reads the module-level variable on every call,
    # so monkeypatching it is sufficient without reloading the module.
    monkeypatch.setattr(analytics_db, 'DATABASE_PATH', db_path)
    analytics_db.init_db()

    # Remove any real API key so no OpenAI calls can escape.
    monkeypatch.delenv('OPENAI_API_KEY', raising=False)

    flask_app.config.update({
        'TESTING': True,
        'WTF_CSRF_ENABLED': False,
        'UPLOAD_FOLDER': upload_dir,
        'REPO_FOLDER': repo_dir,
    })

    # Flask-Limiter reads .enabled as an instance attribute (set once at
    # init_app time), so the RATELIMIT_ENABLED config key has no effect after
    # startup.  Patch the attribute directly instead.
    monkeypatch.setattr(flask_limiter, 'enabled', False)

    yield flask_app

    flask_app.config.update({
        'TESTING': False,
        'WTF_CSRF_ENABLED': True,
    })


@pytest.fixture
def client(app):
    """Flask test client."""
    return app.test_client()


# ---------------------------------------------------------------------------
# OpenAI mock helpers
# ---------------------------------------------------------------------------

def _openai_response(content: str) -> MagicMock:
    """Construct a minimal object that mimics an OpenAI chat completion."""
    resp = MagicMock()
    resp.choices = [MagicMock()]
    resp.choices[0].message.content = content
    return resp


@pytest.fixture
def mock_openai(mocker):
    """
    Patch syllabus_checker.OpenAI so no network calls are made.

    The fixture returns the mock *instance* (client), so tests can configure
    ``mock_openai.chat.completions.create.return_value`` as needed.

    Default response is an empty JSON object ``{}``.
    """
    mock_cls = mocker.patch('syllabus_checker.OpenAI')
    instance = mock_cls.return_value
    instance.chat.completions.create.return_value = _openai_response('{}')
    return instance


@pytest.fixture
def mock_openai_metadata(mock_openai):
    """Pre-configure mock_openai to return a plausible metadata JSON blob."""
    mock_openai.chat.completions.create.return_value = _openai_response(
        '{"instructor_name": "Dr. Test Instructor", '
        '"department": "Computer Science", '
        '"course_subject": "CS", '
        '"course_number": "101", '
        '"course_title": "Introduction to Testing", '
        '"semester": "Fall 2025"}'
    )
    return mock_openai
